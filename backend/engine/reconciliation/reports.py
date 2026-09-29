import csv
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

REPORTS = {
    "ca_reconciliation": "CA Reconciliation Report",
    "exceptions": "Exception Report",
    "control_totals": "Control Totals",
    "party_followup": "Party Follow-up",
    "customer_summary": "Customer Summary",
    "tan_summary": "TAN Summary",
    "quarterly_summary": "Quarterly Summary",
}
FORMATS = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv", "pdf": "application/pdf"}

_RESULT_COLS = [("transaction_id", "Transaction ID"), ("customer", "Customer"), ("customer_code", "Code"), ("party_pan", "PAN"), ("tan", "TAN"), ("deductor_name", "Deductor"), ("books_date", "Books Date"), ("statement_date", "26AS Date"), ("financial_year", "FY"), ("books_quarter", "Books Qtr"), ("statement_quarter", "26AS Qtr"), ("section", "Section"), ("tds_expected", "TDS Expected"), ("tax_deducted", "Tax Deducted"), ("tds_deposited", "TDS Deposited"), ("difference", "Difference"), ("difference_pct", "Diff %"), ("status", "26AS Status"), ("result", "Result"), ("match_method", "Match Method"), ("claimability", "Claimability"), ("identity_status", "Identity"), ("match_group_id", "Group"), ("reason", "Reason"), ("recommended_action", "Recommended Action")]


def _source_refs(rows, key):
    return ", ".join(str(row.get(key) or f"row {row.get('row_no', '?')}") for row in rows or [])


def _sales_report(report, run, rows):
    """Project only the Sales-workflow contract; never coerce it into Books fields."""
    meta = [("Assessee", run.get("assessee_name") or "—"), ("PAN", run.get("assessee_pan") or "—"), ("Financial Year", run.get("financial_year") or "—"), ("Workflow", "Sales + TDS + 26AS"), ("Run ID", run["run_id"])]
    if report == "control_totals":
        s = run.get("summary") or {}
        return REPORTS[report], meta, ["Control", "Value"], [["Sales records", s.get("sales_count")], ["TDS Expected / Receivable records", s.get("tds_count")], ["26AS entries", s.get("statement_count")], ["Amount matched", s.get("amount_matched_count")], ["Amount differences", s.get("amount_difference_count")], ["TDS matched", s.get("tds_matched_count")], ["TDS differences", s.get("tds_difference_count")]]
    cols = ["Record", "Sales Customer", "Customer PAN", "Customer GSTIN", "Deductor", "TAN", "Identity Status", "Identity Method", "Identity Evidence", "Sales Amount", "26AS Amount Credited/Paid", "Amount Difference", "Amount Status", "Sales Source Rows", "26AS Amount Source Rows", "Amount Group", "TDS Expected / Receivable", "26AS TDS Deducted", "TDS Difference", "TDS Status", "TDS Source Rows", "26AS TDS Source Rows", "TDS Group", "Overall Status", "Reason", "Recommended Action"]
    data = []
    for row in rows:
        identity, amount, tds = row.get("identity") or {}, row.get("amount_check") or {}, row.get("tds_check") or {}
        exceptional = identity.get("status") != "CONFIRMED" or amount.get("status") != "AMOUNT_MATCHED" or tds.get("status") != "TDS_MATCHED"
        if report == "exceptions" and not exceptional:
            continue
        data.append([row.get("transaction_id"), identity.get("customer_name"), identity.get("customer_pan"), identity.get("customer_gstin"), identity.get("deductor_name"), identity.get("tan"), identity.get("status"), identity.get("method"), identity.get("reason"), amount.get("sales_amount"), amount.get("statement_amount_paid"), amount.get("difference"), amount.get("status"), _source_refs(amount.get("sales_entries"), "sales_transaction_id"), _source_refs(amount.get("statement_entries"), "statement_id"), amount.get("match_group_id"), tds.get("tds_expected"), tds.get("statement_tds_deducted"), tds.get("difference"), tds.get("status"), _source_refs(tds.get("tds_entries"), "tds_transaction_id"), _source_refs(tds.get("statement_entries"), "statement_id"), tds.get("match_group_id"), row.get("overall_status"), row.get("reason"), row.get("recommended_action")])
    return REPORTS[report], meta, cols, data


def build_report(report: str, run: dict, rows: list, exception_states: dict):
    if run.get("workflow") == "SALES_TDS_26AS":
        return _sales_report(report, run, rows)
    title = REPORTS[report]
    meta = [("Assessee", run.get("assessee_name") or "—"), ("PAN", run.get("assessee_pan") or "—"), ("Financial Year", run.get("financial_year") or "—"), ("Run ID", run["run_id"]), ("Generated", run.get("finished_at") or "")]
    if report == "ca_reconciliation":
        return title, meta, [c[1] for c in _RESULT_COLS], [[r.get(k) for k, _ in _RESULT_COLS] for r in rows]
    if report == "exceptions":
        cols = ["Severity", "Transaction ID", "Customer", "TAN", "Quarter", "Amount", "Difference", "Result", "Reason", "Recommended Action", "Status", "Notes"]
        data = []
        for r in rows:
            if r["result"] == "MATCHED_CLAIMABLE":
                continue
            st = exception_states.get(r["id"], {})
            data.append([r["severity"], r["transaction_id"], r["customer"], r["tan"], r["books_quarter"] or r["statement_quarter"], r["tds_expected"] if r["tds_expected"] is not None else r["tax_deducted"], r["difference"], r["result"], r["reason"], r["recommended_action"], st.get("status", "OPEN"), " | ".join(n["text"] for n in st.get("notes", []))])
        return title, meta, cols, data
    s = run.get("summary") or {}
    if report == "control_totals":
        data = [["Books TDS Expected", s.get("books_tds_expected")], ["26AS Tax Deducted", s.get("statement_tax_deducted")], ["26AS TDS Deposited", s.get("statement_tds_deposited")], ["Matched (Books)", s.get("matched_amount")], ["Matched & Claimable", s.get("matched_claimable_amount")], ["Matched, Not Claimable", s.get("matched_not_claimable_amount")], ["Difference (Books − 26AS)", s.get("difference")], ["Reconciliation %", s.get("reconciliation_percentage")], ["26AS Coverage %", s.get("coverage_percentage")], ["Books rows", s.get("books_count")], ["26AS rows", s.get("statement_count")]]
        data += [[f"Result · {k}", v] for k, v in (s.get("result_counts") or {}).items()]
        data += [[f"Method · {k}", v] for k, v in (s.get("match_method_counts") or {}).items()]
        return title, meta, ["Control", "Value"], data
    if report == "party_followup":
        cols = ["Customer", "Code", "PAN", "TAN(s)", "Transaction ID", "Quarter", "TDS Expected", "Tax Deducted", "Difference", "Issue", "Action"]
        data = [[r["customer"] or r["deductor_name"], r["customer_code"], r["party_pan"], ", ".join(r["matched_tans"]) or r["tan"], r["transaction_id"], r["books_quarter"] or r["statement_quarter"], r["tds_expected"], r["tax_deducted"], r["difference"], r["result"], r["recommended_action"]] for r in rows if r["result"] in ("MISSING_IN_26AS", "AMOUNT_MISMATCH", "MATCHED_NOT_CLAIMABLE", "IDENTITY_UNMAPPED")]
        return title, meta, cols, data
    if report == "customer_summary":
        agg = defaultdict(lambda: {"name": "", "pan": "", "tans": set(), "books": 0, "stmt": 0, "matched": 0, "exc": 0})
        for r in rows:
            a = agg[r["customer_code"] or "(unmapped)"]
            a["name"] = a["name"] or r["customer"] or r["deductor_name"]
            a["pan"] = a["pan"] or r["party_pan"]
            a["tans"].update(r["matched_tans"])
            a["books"] += r["tds_expected"] or 0
            a["stmt"] += (r["tax_deducted"] or 0) if (r["source"] == "26AS" or r["group_size"] == 2 or r["result"] == "AMOUNT_MISMATCH") else 0
            a["matched" if r["result"] in ("MATCHED_CLAIMABLE", "MATCHED_NOT_CLAIMABLE") else "exc"] += 1
        return title, meta, ["Code", "Customer", "PAN", "TANs", "Books TDS", "26AS TDS", "Difference", "Matched rows", "Exception rows"], [[c, a["name"], a["pan"], ", ".join(sorted(a["tans"])), round(a["books"], 2), round(a["stmt"], 2), round(a["books"] - a["stmt"], 2), a["matched"], a["exc"]] for c, a in sorted(agg.items())]
    if report == "tan_summary":
        agg = defaultdict(lambda: {"deductor": "", "customer": "", "code": "", "identity": "", "entries": 0, "deducted": 0, "deposited": 0, "matched": 0})
        for r in rows:
            for e in r["statement_entries"]:
                a = agg[e["tan"]]
                a.update({"deductor": e["deductor_name"], "customer": r["customer"], "code": r["customer_code"], "identity": r["identity_status"]})
                if r["source"] == "26AS" or r["matched_books_ids"][:1] == [r["transaction_id"]]:
                    a["entries"] += 1
                    a["deducted"] += e["tax_deducted"] or 0
                    a["deposited"] += e["tds_deposited"] or 0
                    a["matched"] += 1 if r["result"] in ("MATCHED_CLAIMABLE", "MATCHED_NOT_CLAIMABLE") else 0
        return title, meta, ["TAN", "Deductor", "Customer", "Code", "Identity", "Entries", "Tax Deducted", "TDS Deposited", "Matched entries"], [[t, a["deductor"], a["customer"], a["code"], a["identity"], a["entries"], round(a["deducted"], 2), round(a["deposited"], 2), a["matched"]] for t, a in sorted(agg.items())]
    if report == "quarterly_summary":
        return title, meta, ["Quarter", "Books TDS", "26AS TDS", "Matched rows", "Exception rows"], [[q["quarter"], q["books_tds"], q["statement_tds"], q["matched"], q["exceptions"]] for q in s.get("by_quarter", [])]
    raise KeyError(report)


def render(fmt: str, title: str, meta: list, cols: list, data: list) -> bytes:
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        for k, v in meta:
            w.writerow([k, v])
        w.writerow([])
        w.writerow(cols)
        w.writerows(data)
        return buf.getvalue().encode("utf-8-sig")
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = title[:30]
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        for k, v in meta:
            ws.append([k, v])
        ws.append([])
        ws.append(cols)
        hdr = ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=hdr, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="10233F")
            cell.alignment = Alignment(vertical="center")
        for row in data:
            ws.append(row)
        for i, col in enumerate(cols, start=1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in data[:200] if r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    if fmt == "pdf":
        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
        styles = getSampleStyleSheet()
        small = styles["BodyText"]
        small.fontSize = 6.5
        small.leading = 8
        hdr_style = ParagraphStyle("hdr", parent=small, textColor=colors.white)
        story = [Paragraph(title, styles["Title"]), Paragraph(" · ".join(f"<b>{k}:</b> {v}" for k, v in meta), styles["BodyText"]), Spacer(1, 6)]
        table_data = [[Paragraph(f"<b>{c}</b>", hdr_style) for c in cols]] + [[Paragraph(_pdf_cell(v), small) for v in row] for row in data]
        avail = landscape(A4)[0] - 20 * mm
        widths = [avail / len(cols)] * len(cols)
        t = Table(table_data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10233F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D3DF")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")])]))
        story.append(t)
        doc.build(story)
        return out.getvalue()
    raise KeyError(fmt)


def _pdf_cell(v):
    if v is None:
        return ""
    return str(v).replace("&", "&amp;").replace("<", "&lt;")
